"""
Opening-stock import.

Reads a CSV or Excel file, validates EVERY row, and only then posts stock
through record_stock_movement() so the audit trail stays intact (each line
becomes an ADJUSTMENT_IN movement with reference OPENING-STOCK).

All-or-nothing: if any row has a problem, nothing is imported.
"""

import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction

from .models import Product, StockMovement, Warehouse
from .services.stock import record_stock_movement

OPENING_REF = "OPENING-STOCK"

COLUMNS = ["sku", "name", "category", "unit", "cost_price", "selling_price",
           "reorder_level", "warehouse", "quantity"]


def _norm_header(h):
    return str(h or "").strip().lower().replace(" ", "_")


def parse_file(uploaded_file):
    """Returns a list of dicts keyed by normalised header. Raises ValueError for unreadable files."""
    filename = (uploaded_file.name or "").lower()
    raw = uploaded_file.read()

    if filename.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        table = list(reader)
    elif filename.endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        except Exception:
            raise ValueError("Could not open that Excel file. Make sure it is a real .xlsx file.")
        ws = wb.worksheets[0]
        table = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        raise ValueError("Please upload a .xlsx or .csv file.")

    if not table:
        raise ValueError("The file is empty.")

    headers = [_norm_header(h) for h in table[0]]
    if "sku" not in headers or "warehouse" not in headers or "quantity" not in headers:
        raise ValueError("The first row must contain the headings: sku, warehouse, quantity "
                         "(see the template for all columns).")

    rows = []
    for i, cells in enumerate(table[1:], start=2):
        record = {}
        for h, c in zip(headers, cells):
            if h:
                record[h] = "" if c is None else str(c).strip()
        if not any(record.values()):
            continue  # skip fully blank rows
        record["_row"] = i
        rows.append(record)

    if not rows:
        raise ValueError("No data rows found under the headings.")
    return rows


def _to_int(value):
    s = str(value).strip().replace(",", "")
    if s == "":
        raise ValueError("missing")
    d = Decimal(s)
    if d != d.to_integral_value():
        raise ValueError("not a whole number")
    return int(d)


def _to_money(value):
    s = str(value).strip().replace(",", "").replace("₦", "").replace("NGN", "")
    if s == "":
        raise ValueError("missing")
    d = Decimal(s)
    if d < 0:
        raise ValueError("negative")
    return d


def _match_unit(value):
    v = (value or "").strip().lower()
    if not v:
        return Product.Unit.PIECE
    for code, label in Product.Unit.choices:
        if v in (code.lower(), label.lower()):
            return code
    return None


def validate_rows(rows):
    """Returns (clean_rows, errors). errors is a list of 'Row N: message' strings."""
    errors = []
    clean = []
    seen = set()
    warehouses = {w.name.strip().lower(): w for w in Warehouse.objects.all()}
    wh_names = ", ".join(sorted(w.name for w in warehouses.values())) or "none created yet"

    # Details for a new product may be written on just one of its rows
    # (e.g. only the first warehouse row). Share them across all rows of that SKU.
    merged = {}
    for r in rows:
        m = merged.setdefault(r.get("sku", "").strip().lower(), {})
        for f in ("name", "category", "unit", "cost_price", "selling_price", "reorder_level"):
            value = str(r.get(f, "") if r.get(f, "") is not None else "").strip()
            if not m.get(f) and value:
                m[f] = value

    for r in rows:
        n = r["_row"]
        sku = r.get("sku", "").strip()
        d = merged.get(sku.lower(), {})
        wh_name = r.get("warehouse", "").strip()
        problems = []

        if not sku:
            problems.append("sku is missing")
        if not wh_name:
            problems.append("warehouse is missing")

        wh = warehouses.get(wh_name.lower()) if wh_name else None
        if wh_name and wh is None:
            problems.append(f"warehouse '{wh_name}' not found (existing: {wh_names})")

        try:
            qty = _to_int(r.get("quantity", ""))
            if qty < 0:
                problems.append("quantity cannot be negative")
        except (ValueError, InvalidOperation):
            qty = None
            problems.append(f"quantity '{r.get('quantity', '')}' is not a valid whole number")

        product = Product.objects.filter(sku__iexact=sku).first() if sku else None
        new_product = None
        if sku and product is None:
            name = d.get("name", "")
            if not name:
                missing = [f for f in ("name", "cost_price", "selling_price") if not d.get(f)]
                problems.append(f"sku '{sku}' is not in the system. To create it, fill in: {', '.join(missing)}")
            else:
                try:
                    cost = _to_money(d.get("cost_price", ""))
                    sell = _to_money(d.get("selling_price", ""))
                except (ValueError, InvalidOperation):
                    problems.append(f"new product '{sku}' needs valid numbers for cost_price and selling_price")
                    cost = sell = None
                unit = _match_unit(d.get("unit", ""))
                if unit is None:
                    problems.append(f"unit '{d.get('unit')}' not recognised (use Bag, Carton, Piece or Roll)")
                try:
                    reorder = _to_int(d.get("reorder_level", "") or "0")
                except (ValueError, InvalidOperation):
                    reorder = 0
                    problems.append("reorder_level must be a whole number")
                if cost is not None and sell is not None and unit is not None:
                    new_product = {
                        "sku": sku.upper(), "name": name, "category": d.get("category", ""),
                        "unit": unit, "cost_price": str(cost), "selling_price": str(sell),
                        "reorder_level": reorder,
                    }

        if product is not None and wh is not None:
            already = StockMovement.objects.filter(
                product=product, warehouse=wh, reference=OPENING_REF).exists()
            if already:
                problems.append(f"opening stock was already loaded for '{sku}' at {wh.name} (importing again would double it)")

        key = (sku.lower(), wh_name.lower())
        if sku and wh_name:
            if key in seen:
                problems.append(f"'{sku}' at '{wh_name}' appears more than once in the file")
            seen.add(key)

        if problems:
            for p in problems:
                errors.append(f"Row {n}: {p}")
            continue

        clean.append({
            "row": n, "sku": sku, "warehouse_id": wh.id, "warehouse_name": wh.name,
            "quantity": qty, "new_product": new_product,
            "product_name": product.name if product else new_product["name"],
        })

    return clean, errors


@transaction.atomic
def apply_import(clean_rows, user=None):
    """Posts the validated rows. Returns a small summary dict."""
    created_products = 0
    lines = 0
    units = 0
    for r in clean_rows:
        product = Product.objects.filter(sku__iexact=r["sku"]).first()
        if product is None:
            np = r["new_product"]
            product = Product.objects.create(
                sku=np["sku"], name=np["name"], category=np["category"], unit=np["unit"],
                cost_price=Decimal(np["cost_price"]), selling_price=Decimal(np["selling_price"]),
                reorder_level=np["reorder_level"],
            )
            created_products += 1
        warehouse = Warehouse.objects.get(pk=r["warehouse_id"])
        if r["quantity"] > 0:
            record_stock_movement(
                product=product, warehouse=warehouse, quantity=r["quantity"],
                movement_type=StockMovement.MovementType.ADJUSTMENT_IN,
                reference=OPENING_REF, user=user, notes="Opening stock import",
            )
            lines += 1
            units += r["quantity"]
    return {"products_created": created_products, "stock_lines": lines, "units": units}


def build_template_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Opening stock"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2D4A3E")
    widths = [16, 30, 14, 10, 12, 14, 14, 20, 10]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w

    guide = wb.create_sheet("How to fill this in")
    for line in [
        "One row per product PER WAREHOUSE. Same product in two warehouses = two rows.",
        "Required on every row: sku, warehouse, quantity.",
        "warehouse must match a warehouse name already in the system exactly (capitals do not matter).",
        "If the sku is already in the system, the other columns are ignored.",
        "If the sku is NEW, also fill name, cost_price and selling_price (unit, category, reorder_level are optional).",
        "unit can be: Bag, Carton, Piece or Roll. Blank = Piece.",
        "Quantity is the count you physically have TODAY.",
        "Load opening stock only ONCE per product and warehouse. The system blocks a second load.",
        "Delete this sheet or leave it - only the first sheet is read.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
